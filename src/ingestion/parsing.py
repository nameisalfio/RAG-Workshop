"""Parsing stage: turn a raw file of any supported type into text.

PDF files go to one of several interchangeable parsers; text, DOCX and CSV files
have a dedicated reader; every other type goes to Docling.
"""

from __future__ import annotations

import csv
import io
import json
import logging
from bisect import bisect_right
from collections.abc import Sequence
from dataclasses import dataclass
from itertools import accumulate
from pathlib import Path
from typing import Protocol, cast

import hydra
import pdfplumber
import pymupdf
import pymupdf4llm
from docx import Document
from docx.text.paragraph import Paragraph
from omegaconf import DictConfig, OmegaConf
from pypdf import PdfReader

from src.config import CONFIG_NAME, ParsingConfig, RootConfig, project_path, register_schema

LOGGER = logging.getLogger(__name__)

PDF_BACKENDS = ("pypdf", "pymupdf", "pdfplumber", "pymupdf4llm", "docling")
GENERIC_BACKENDS = ("docling",)
_TEXT_EXTENSIONS = frozenset({".txt", ".md"})
_PAGE_SEPARATOR = "\n\n"
_PARSED_SUFFIX = ".json"


class MissingTextLayerError(RuntimeError):
    """Raised when a page yields no text, typically a scan that needs OCR."""


class UnsupportedFormatError(ValueError):
    """Raised when no parser is available for a file type."""


@dataclass(frozen=True)
class ParsedDocument:
    """Text of one file, one entry per page.

    ``paginated`` is False for formats without physical pages (text, DOCX, CSV):
    their content is a single entry and its position is not a page citation.
    """

    path: Path
    parser: str
    pages: tuple[str, ...]
    paginated: bool

    @property
    def text(self) -> str:
        """Text of the whole document, pages separated by a blank line."""
        return _PAGE_SEPARATOR.join(self.pages)

    def page_of(self, offset: int) -> int | None:
        """Page number, from 1, that holds a character offset of ``text``; None for formats without pages."""
        if not self.paginated:
            return None
        starts = list(accumulate((len(page) + len(_PAGE_SEPARATOR) for page in self.pages[:-1]), initial=0))
        return bisect_right(starts, offset)


class Parser(Protocol):
    """Extracts the text of a file."""

    name: str

    def parse(self, path: Path) -> ParsedDocument:
        """Parse a file; a page without text raises MissingTextLayerError."""
        ...


def _document(path: Path, parser: str, pages: Sequence[str], paginated: bool) -> ParsedDocument:
    """Assemble a parsed document, refusing empty documents and pages without text."""
    if not pages:
        raise ValueError(f"{path}: parser {parser} returned no content")
    for number, text in enumerate(pages, start=1):
        if not text.strip():
            raise MissingTextLayerError(f"{path}: parser {parser} found no text on page {number}")
    return ParsedDocument(path, parser, tuple(text.strip() for text in pages), paginated)


def list_documents(root: Path) -> list[Path]:
    """Return every non-hidden file below a directory, which must exist and hold at least one file."""
    if not root.is_dir():
        raise FileNotFoundError(f"Documents directory not found: {root}")
    files = sorted(path for path in root.rglob("*") if path.is_file() and not path.name.startswith("."))
    if not files:
        raise FileNotFoundError(f"Documents directory is empty: {root}")
    return files


# ------------------------------------------------------------------ PDF parsers


class PypdfParser:
    """Pure-Python text extraction, in the order of the PDF content stream."""

    name = "pypdf"

    def parse(self, path: Path) -> ParsedDocument:
        """Extract the text of each page with pypdf."""
        return _document(path, self.name, [page.extract_text() for page in PdfReader(path).pages], True)


class PymupdfParser:
    """Text extraction with PyMuPDF, block by block."""

    name = "pymupdf"

    def parse(self, path: Path) -> ParsedDocument:
        """Extract the text of each page with PyMuPDF."""
        with pymupdf.open(path) as pdf:  # type: ignore[no-untyped-call]
            return _document(path, self.name, [page.get_text("text") for page in pdf], True)


class PdfplumberParser:
    """Text extraction with pdfplumber, rebuilt from character positions."""

    name = "pdfplumber"

    def parse(self, path: Path) -> ParsedDocument:
        """Extract the text of each page with pdfplumber."""
        with pdfplumber.open(path) as pdf:
            return _document(path, self.name, [page.extract_text() for page in pdf.pages], True)


class Pymupdf4llmParser:
    """Markdown extraction with pymupdf4llm: headings and tables are kept as Markdown."""

    name = "pymupdf4llm"

    def parse(self, path: Path) -> ParsedDocument:
        """Convert each page to Markdown with pymupdf4llm."""
        chunks = pymupdf4llm.to_markdown(str(path), page_chunks=True, show_progress=False)
        return _document(path, self.name, [chunk["text"] for chunk in chunks], True)


class DoclingParser:
    """Markdown extraction with Docling: layout analysis, table structure, OCR and many input formats."""

    name = "docling"

    def __init__(self) -> None:
        # Imported on demand: Docling loads a deep-learning stack that the other parsers do not need.
        from docling.document_converter import DocumentConverter

        self._converter = DocumentConverter()

    def parse(self, path: Path) -> ParsedDocument:
        """Convert a file to Markdown, page by page when the format has pages."""
        document = self._converter.convert(path).document
        numbers = sorted(document.pages)
        if not numbers:
            return _document(path, self.name, [document.export_to_markdown()], False)
        return _document(path, self.name, [document.export_to_markdown(page_no=n) for n in numbers], True)


def docling_extensions() -> frozenset[str]:
    """Return the file extensions Docling can convert, lowercase and with a leading dot."""
    # Imported on demand for the same reason as in DoclingParser.
    from docling.datamodel.base_models import FormatToExtensions

    return frozenset(f".{extension.lower()}" for group in FormatToExtensions.values() for extension in group)


def build_pdf_parser(backend: str) -> Parser:
    """Create the PDF parser with the given name."""
    parsers: dict[str, type[Parser]] = {
        "pypdf": PypdfParser,
        "pymupdf": PymupdfParser,
        "pdfplumber": PdfplumberParser,
        "pymupdf4llm": Pymupdf4llmParser,
        "docling": DoclingParser,
    }
    if backend not in parsers:
        raise ValueError(f"Unknown PDF backend {backend!r}; expected one of {PDF_BACKENDS}")
    return parsers[backend]()


# ---------------------------------------------------- text, DOCX and CSV readers


def markdown_table(header: Sequence[str], rows: Sequence[Sequence[str]]) -> str:
    """Render a table as Markdown; every row must have as many cells as the header."""

    def line(cells: Sequence[str]) -> str:
        if len(cells) != len(header):
            raise ValueError(f"Row {list(cells)!r} has {len(cells)} cells, expected {len(header)}")
        return "| " + " | ".join(" ".join(cell.split()).replace("|", "\\|") for cell in cells) + " |"

    separator = "| " + " | ".join("---" for _ in header) + " |"
    return "\n".join([line(header), separator, *(line(row) for row in rows)])


def read_text(path: Path, encoding: str) -> ParsedDocument:
    """Read a plain-text or Markdown file."""
    return _document(path, "text", [path.read_text(encoding=encoding)], False)


def read_docx(path: Path) -> ParsedDocument:
    """Read a DOCX file as Markdown, keeping headings, lists and tables in document order."""
    blocks: list[str] = []
    for item in Document(str(path)).iter_inner_content():
        if not isinstance(item, Paragraph):
            rows = [[cell.text for cell in row.cells] for row in item.rows]
            blocks.append(markdown_table(rows[0], rows[1:]))
            continue
        text = item.text.strip()
        style = item.style.name if item.style is not None and item.style.name else ""
        level = style.removeprefix("Heading ")
        if not text:
            continue
        if style == "Title":
            blocks.append(f"# {text}")
        elif style.startswith("Heading ") and level.isdigit():
            blocks.append(f"{'#' * (int(level) + 1)} {text}")
        elif style.startswith("List"):
            blocks.append(f"- {text}")
        else:
            blocks.append(text)
    return _document(path, "python-docx", ["\n\n".join(blocks)], False)


def read_csv(path: Path, encoding: str, delimiters: str) -> ParsedDocument:
    """Read a CSV file as one Markdown table.

    The delimiter is the candidate that occurs most often in the header line.
    """
    content = path.read_text(encoding=encoding)
    header_line = next((line for line in content.splitlines() if line.strip()), "")
    counts = sorted(((header_line.count(candidate), candidate) for candidate in delimiters), reverse=True)
    if counts[0][0] == 0 or (len(counts) > 1 and counts[0][0] == counts[1][0]):
        raise ValueError(f"{path}: cannot detect the delimiter among {delimiters!r} from the header line")
    rows = [row for row in csv.reader(io.StringIO(content), delimiter=counts[0][1]) if row]
    if len(rows) < 2:
        raise ValueError(f"{path}: a CSV file needs a header and at least one data row")
    try:
        return _document(path, "csv", [markdown_table(rows[0], rows[1:])], False)
    except ValueError as error:
        raise ValueError(f"{path}: {error}") from error


# ----------------------------------------------------------------------- routing


class DocumentRouter:
    """Sends each file to the parser for its type.

    Types without a dedicated reader go to the generic backend, created on first
    use because it is slow to start; a type it cannot read is an error.
    """

    def __init__(self, config: ParsingConfig) -> None:
        if config.generic_backend not in GENERIC_BACKENDS:
            raise ValueError(
                f"Unknown generic backend {config.generic_backend!r}; expected one of {GENERIC_BACKENDS}"
            )
        self._config = config
        self._pdf_parser = build_pdf_parser(config.pdf_backend)
        # One Docling instance serves both roles: its models are loaded only once.
        self._generic_parser = self._pdf_parser if isinstance(self._pdf_parser, DoclingParser) else None

    def parse(self, path: Path) -> ParsedDocument:
        """Parse a file according to its extension."""
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._pdf_parser.parse(path)
        if suffix in _TEXT_EXTENSIONS:
            return read_text(path, self._config.text_encoding)
        if suffix == ".docx":
            return read_docx(path)
        if suffix == ".csv":
            return read_csv(path, self._config.text_encoding, self._config.csv_delimiters)
        if suffix not in docling_extensions():
            raise UnsupportedFormatError(f"No parser can read files of type {suffix!r}: {path}")
        if self._generic_parser is None:
            self._generic_parser = DoclingParser()
        return self._generic_parser.parse(path)


# ------------------------------------------------------------------- persistence


def save_parsed(document: ParsedDocument, path: Path) -> None:
    """Write a parsed document as JSON."""
    record = {
        "source": document.path.as_posix(),
        "parser": document.parser,
        "paginated": document.paginated,
        "pages": list(document.pages),
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")


def load_parsed(path: Path) -> ParsedDocument:
    """Read a parsed document written by save_parsed."""
    record = json.loads(path.read_text(encoding="utf-8"))
    return ParsedDocument(
        Path(record["source"]), record["parser"], tuple(record["pages"]), record["paginated"]
    )


def parse_corpus(raw_dir: Path, parsed_dir: Path, router: DocumentRouter) -> list[Path]:
    """Parse every raw document and save it under the same relative path; the first failure stops the run.

    Saved documents record their path relative to the raw directory, so that a
    chunk can cite its source wherever the corpus is stored.
    """
    saved: list[Path] = []
    for path in list_documents(raw_dir):
        source = path.relative_to(raw_dir)
        parsed = router.parse(path)
        target = parsed_dir / f"{source.as_posix()}{_PARSED_SUFFIX}"
        save_parsed(ParsedDocument(source, parsed.parser, parsed.pages, parsed.paginated), target)
        LOGGER.info("Parsed %s with %s (%d pages)", source, parsed.parser, len(parsed.pages))
        saved.append(target)
    return saved


def load_corpus(parsed_dir: Path) -> list[ParsedDocument]:
    """Read every parsed document saved by parse_corpus."""
    return [load_parsed(path) for path in list_documents(parsed_dir) if path.suffix == _PARSED_SUFFIX]


# ------------------------------------------------------------------ entry point


@hydra.main(version_base=None, config_path="../conf", config_name=CONFIG_NAME)
def main(cfg: DictConfig) -> None:
    """Parse every raw document and save the text; run it with ``python -m src.ingestion.parsing``."""
    config = cast(RootConfig, OmegaConf.to_object(cfg))
    raw_dir = project_path(config.paths.raw_documents_dir)
    parsed_dir = project_path(config.paths.parsed_dir)
    saved = parse_corpus(raw_dir, parsed_dir, DocumentRouter(config.parsing))
    LOGGER.info("Parsed %d documents into %s", len(saved), parsed_dir)


if __name__ == "__main__":
    register_schema()
    main()
