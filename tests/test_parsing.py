"""Parsers must return texts and refuse what they cannot read."""

from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest
from docx import Document

from src.config import ParsingConfig
from src.ingestion.parsing import (
    DocumentRouter,
    MissingTextLayerError,
    ParsedDocument,
    UnsupportedFormatError,
    build_pdf_parser,
    list_documents,
    load_corpus,
    parse_corpus,
)

LIGHT_BACKENDS = ("pypdf", "pymupdf", "pdfplumber", "pymupdf4llm")
SENTENCE = "The withdrawal notice is 45 days."


def make_router(pdf_backend: str = "pymupdf") -> DocumentRouter:
    """Build a router that never needs the generic backend unless a test asks for an unusual type."""
    return DocumentRouter(
        ParsingConfig(
            pdf_backend=pdf_backend, generic_backend="docling", text_encoding="utf-8", csv_delimiters=",;\t|"
        )
    )


def write_pdf(path: Path, pages: list[str]) -> None:
    """Write a PDF with one line of text per page; an empty string gives a page without text."""
    with pymupdf.open() as pdf:  # type: ignore[no-untyped-call]
        for text in pages:
            page = pdf.new_page()
            if text:
                page.insert_text((72, 72), text)
        pdf.save(path)


@pytest.mark.parametrize("backend", LIGHT_BACKENDS)
def test_pdf_backends_return_one_text_per_page(tmp_path: Path, backend: str) -> None:
    path = tmp_path / "doc.pdf"
    write_pdf(path, [SENTENCE, "Second page."])
    document = make_router(backend).parse(path)
    assert document.parser == backend
    assert document.paginated
    assert len(document.pages) == 2
    assert SENTENCE in document.pages[0]


@pytest.mark.parametrize("backend", LIGHT_BACKENDS)
def test_page_without_text_is_an_error(tmp_path: Path, backend: str) -> None:
    path = tmp_path / "scan.pdf"
    write_pdf(path, [""])
    with pytest.raises(MissingTextLayerError):
        build_pdf_parser(backend).parse(path)


def test_text_file_is_one_unpaginated_text(tmp_path: Path) -> None:
    note = tmp_path / "note.md"
    note.write_text("# Title\n\nBody.", encoding="utf-8")
    document = make_router().parse(note)
    assert document.text == "# Title\n\nBody."
    assert not document.paginated


def test_docx_keeps_headings_lists_and_tables_in_order(tmp_path: Path) -> None:
    path = tmp_path / "letter.docx"
    source = Document()
    source.add_heading("Withdrawal letter", level=0)
    source.add_heading("Recipient", level=1)
    source.add_paragraph("Sent by certified e-mail.")
    source.add_paragraph("Attach the contract", style="List Bullet")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Notice"
    table.cell(1, 1).text = "45 days"
    source.save(str(path))
    text = make_router().parse(path).text
    assert text.splitlines()[0] == "# Withdrawal letter"
    assert "## Recipient" in text
    assert "- Attach the contract" in text
    assert "| Notice | 45 days |" in text
    assert text.index("Sent by certified e-mail.") < text.index("| Field | Value |")


@pytest.mark.parametrize("delimiter", [",", ";", "\t"])
def test_csv_delimiter_is_detected(tmp_path: Path, delimiter: str) -> None:
    path = tmp_path / "tickets.csv"
    rows = [
        ["ticket", "code", "hours"],
        ["TK-1", "E-4012", "3,5" if delimiter != "," else "3.5"],
        ["TK-2", "E-4021", "2"],
    ]
    path.write_text("\n".join(delimiter.join(row) for row in rows) + "\n", encoding="utf-8")
    text = make_router().parse(path).text
    assert text.splitlines()[0] == "| ticket | code | hours |"
    assert "| TK-2 | E-4021 | 2 |" in text


def test_csv_with_ragged_rows_or_no_data_is_an_error(tmp_path: Path) -> None:
    ragged = tmp_path / "ragged.csv"
    ragged.write_text("a;b;c\n1;2;3\n4;5\n", encoding="utf-8")
    with pytest.raises(ValueError, match="cells"):
        make_router().parse(ragged)
    header_only = tmp_path / "header.csv"
    header_only.write_text("a;b;c\n", encoding="utf-8")
    with pytest.raises(ValueError):
        make_router().parse(header_only)


def test_page_of_maps_text_offsets_to_pages() -> None:
    document = ParsedDocument(Path("doc.pdf"), "test", ("first page", "second"), paginated=True)
    assert document.text == "first page\n\nsecond"
    assert document.page_of(0) == 1
    assert document.page_of(document.text.index("second")) == 2
    assert ParsedDocument(Path("note.md"), "test", ("body",), paginated=False).page_of(0) is None


def test_corpus_is_saved_with_relative_sources_and_read_back(tmp_path: Path) -> None:
    raw_dir, parsed_dir = tmp_path / "raw", tmp_path / "parsed"
    (raw_dir / "contracts").mkdir(parents=True)
    write_pdf(raw_dir / "contracts" / "doc.pdf", [SENTENCE, "Second page."])
    (raw_dir / "note.txt").write_text("A plain note.", encoding="utf-8")
    saved = parse_corpus(raw_dir, parsed_dir, make_router())
    assert [path.relative_to(parsed_dir).as_posix() for path in saved] == [
        "contracts/doc.pdf.json",
        "note.txt.json",
    ]
    documents = load_corpus(parsed_dir)
    assert [document.path.as_posix() for document in documents] == ["contracts/doc.pdf", "note.txt"]
    assert documents[0].paginated and len(documents[0].pages) == 2
    assert documents[1].text == "A plain note."


def test_unknown_type_backend_and_directory_are_errors(tmp_path: Path) -> None:
    other = tmp_path / "data.unknownformat"
    other.write_bytes(b"binary")
    with pytest.raises(UnsupportedFormatError):
        make_router().parse(other)
    with pytest.raises(ValueError, match="Unknown PDF backend"):
        build_pdf_parser("magic")
    with pytest.raises(FileNotFoundError):
        list_documents(tmp_path / "missing")
    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(FileNotFoundError, match="empty"):
        list_documents(empty)
